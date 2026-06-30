"""
GST Sales Report Page for the Accounting Desktop Application.
Provides B2B, B2CL, and B2CS classification for GST compliance.
"""
from PySide6.QtWidgets import QWidget, QVBoxLayout, QHBoxLayout, QGridLayout, QLabel, QLineEdit, QDateEdit, QPushButton, QFrame, QTableWidget, QTableWidgetItem, QAbstractItemView, QDialog, QMessageBox, QSizePolicy, QTabWidget, QFileDialog
from PySide6.QtCore import Qt, QDate, QObject, QThread, Signal
from PySide6.QtGui import QColor
from config import active_company_manager
from db import Database
from bizora_core.gst_compliance import classify_invoice, is_interstate, normalized_gstin, place_of_supply_label, tax_totals_for_supply
from bizora_core.gstr1_logic import GSTR1Logic
from ui.table_header_utils import apply_adjustable_table_columns, apply_read_only_report_table_selection, apply_compact_report_table_rows
from ui import theme
from ui.book_report_common import compact_date_style, page_background_style, report_filter_frame_style, report_detail_dialog_style, report_compound_entry_page_style, report_page_shell_style, page_heading_style, report_status_label_style, report_data_table_style, footer_title_style, footer_value_style, report_summary_label_style
from ui.date_formats import configure_qdate_edit, format_display_date, prepare_report_date_edit, qdate_to_db, qdate_to_display
from ui.ui_memory import UiMemoryMixin

PDF_AVAILABLE = False
try:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

class GSTSalesReportWorker(QObject):
    """Load and classify GST sales report data outside the GUI thread."""
    data_ready = Signal(dict)
    error = Signal(str)
    finished = Signal()

    def __init__(self, db_type, db_path, company_id, company_state, from_date, to_date, search):
        super().__init__()
        self.db_type = db_type
        self.db_path = db_path
        self.company_id = company_id
        self.company_state = company_state
        self.from_date = from_date
        self.to_date = to_date
        self.search = search

    def run(self):
        """Fetch invoices, classify sections, and aggregate HSN data."""
        worker_db = None
        try:
            worker_db = Database(db_type=self.db_type, db_path=self.db_path)
            sales = worker_db.get_sales_for_gst_report(self.company_id, self.from_date, self.to_date, self.search)
            b2b_sales, b2cl_sales, b2cs_data = self._classify_sales(sales)
            hsn_data = GSTR1Logic(worker_db)._generate_hsn_summary(self.company_id, self.from_date, self.to_date)
            self.data_ready.emit({'b2b': b2b_sales, 'b2cl': b2cl_sales, 'b2cs': b2cs_data, 'hsn': hsn_data})
        except Exception as exc:
            self.error.emit(str(exc))
        finally:
            if worker_db is not None:
                try:
                    worker_db.force_disconnect()
                except Exception:
                    pass
            self.finished.emit()

    def _classify_sales(self, sales):
        """Classify invoices into B2B, B2CL, and B2CS sections."""
        b2b_sales = []
        b2cl_sales = []
        b2cs_data = {}
        home_state = place_of_supply_label(self.company_state)
        for sale in sales:
            self._normalise_sale_tax_values(sale)
            gstin = normalized_gstin(sale.get('party_gstin') or sale.get('gstin'))
            pos = place_of_supply_label(sale.get('party_state') or sale.get('state'), home_state)
            grand_total = self._to_float(sale.get('grand_total'))
            form_of_sale = classify_invoice(gstin, pos, home_state, grand_total)
            sale['derived_form_of_sale'] = form_of_sale
            sale['place_of_supply'] = pos
            sale['invoice_type'] = 'Regular'
            sale['reverse_charge'] = 'N'
            sale['rate'] = self._get_tax_rate(sale)
            sale['type'] = form_of_sale
            if form_of_sale == 'B2B':
                sale['gstin'] = gstin
                b2b_sales.append(sale)
            elif form_of_sale == 'B2CL':
                b2cl_sales.append(sale)
            else:
                key = (sale['place_of_supply'], sale['rate'])
                if key not in b2cs_data:
                    b2cs_data[key] = {'type': 'B2CS', 'place_of_supply': sale['place_of_supply'], 'rate': sale['rate'], 'taxable_value': 0.0, 'cgst': 0.0, 'sgst': 0.0, 'igst': 0.0, 'cess': 0.0, 'total_value': 0.0, 'invoice_count': 0, 'sale_ids': []}
                b2cs_data[key]['taxable_value'] += self._to_float(sale.get('taxable_value'))
                b2cs_data[key]['cgst'] += self._to_float(sale.get('cgst'))
                b2cs_data[key]['sgst'] += self._to_float(sale.get('sgst'))
                b2cs_data[key]['igst'] += self._to_float(sale.get('igst'))
                b2cs_data[key]['cess'] += self._to_float(sale.get('cess'))
                b2cs_data[key]['total_value'] += grand_total
                b2cs_data[key]['invoice_count'] += 1
                b2cs_data[key]['sale_ids'].append(sale.get('id'))
        return (b2b_sales, b2cl_sales, b2cs_data)

    def _normalise_sale_tax_values(self, sale):
        """Map split tax fields according to POS versus home state."""
        taxable = self._to_float(sale.get('taxable_value'))
        tax_total = self._to_float(sale.get('tax_total') or sale.get('item_tax_total'))
        pos = place_of_supply_label(sale.get('party_state') or sale.get('state'), self.company_state)
        interstate = is_interstate(pos, self.company_state)
        igst, cgst, sgst, cess = tax_totals_for_supply(sale, interstate)
        sale['taxable_value'] = taxable
        sale['cgst'] = cgst
        sale['sgst'] = sgst
        sale['igst'] = igst
        sale['cess'] = cess
        sale['tax_total'] = tax_total or igst + cgst + sgst + cess

    def _get_tax_rate(self, sale):
        """Calculate effective tax rate from normalized tax totals."""
        taxable = self._to_float(sale.get('taxable_value'))
        tax = self._to_float(sale.get('cgst')) + self._to_float(sale.get('sgst')) + self._to_float(sale.get('igst')) + self._to_float(sale.get('cess'))
        if taxable > 0:
            return round(tax / taxable * 100, 2)
        return 0.0

    def _to_float(self, value):
        """Convert values used by worker calculations to float."""
        try:
            return float(value or 0)
        except (TypeError, ValueError):
            return 0.0

class GSTSalesReportPage(UiMemoryMixin, QWidget):
    """GST Sales Report page with tabbed B2B/B2CL, B2CS, and HSN views."""
    MONEY_KEYS = {
        'Invoice Value', 'Taxable Value', 'CGST', 'SGST', 'IGST', 'Cess',
        'Total Value', 'Total Quantity', 'Rate (%)',
    }

    def __init__(self, db=None):
        super().__init__()
        self.db = db or Database()
        self.company_id = None
        self.company_state = ''
        self._loading = False
        self._report_thread = None
        self._report_worker = None
        self.current_report_data = {'b2b': [], 'b2cl': [], 'b2cs': {}, 'hsn': []}
        self.setup_ui()
        self.load_company()
        self._init_ui_memory(table_attrs=("b2b_b2cl_table", "b2cs_table", "hsn_table"))
        self._activate_tab_table(self.tab_widget.currentIndex())

    def setup_ui(self):
        self.setStyleSheet(report_compound_entry_page_style())
        layout = QVBoxLayout(self)
        layout.setContentsMargins(10, 10, 10, 10)
        layout.setSpacing(5)
        self.title_label = QLabel('GST Sales Report')
        self.title_label.setStyleSheet(page_heading_style(18))
        layout.addWidget(self.title_label)
        self.filter_frame = QFrame()
        self.filter_frame.setStyleSheet(report_filter_frame_style())
        filter_layout = QHBoxLayout(self.filter_frame)
        filter_layout.setSpacing(10)
        self.from_date = QDateEdit()
        self.from_date.setDate(QDate.currentDate().addDays(-30))
        prepare_report_date_edit(self.from_date, style_sheet=compact_date_style())
        self.to_date = QDateEdit()
        self.to_date.setDate(QDate.currentDate())
        prepare_report_date_edit(self.to_date, style_sheet=compact_date_style())
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText('Invoice No / Party Name')
        for label_text, widget in (
            ('From Date:', self.from_date),
            ('To Date:', self.to_date),
            ('Search:', self.search_input),
        ):
            filter_layout.addWidget(QLabel(label_text))
            filter_layout.addWidget(widget)
        generate_btn = QPushButton('Generate')
        generate_btn.clicked.connect(self.generate_report)
        export_excel_btn = QPushButton('Export Excel')
        export_excel_btn.clicked.connect(self.export_excel)
        export_pdf_btn = QPushButton('Export PDF')
        export_pdf_btn.clicked.connect(self.export_pdf)
        export_pdf_btn.setEnabled(PDF_AVAILABLE)
        self.generate_btn = generate_btn
        self.export_btn = export_excel_btn
        self.export_pdf_btn = export_pdf_btn
        filter_layout.addWidget(generate_btn)
        filter_layout.addWidget(export_excel_btn)
        filter_layout.addWidget(export_pdf_btn)
        filter_layout.addStretch()
        layout.addWidget(self.filter_frame)
        self.status_label = QLabel('Ready')
        self.status_label.setStyleSheet(report_status_label_style())
        layout.addWidget(self.status_label)
        self.tab_widget = QTabWidget()
        self.b2b_b2cl_table = self._create_report_table(self._b2b_b2cl_headers())
        self.b2cs_table = self._create_report_table(self._b2cs_headers())
        self.hsn_table = self._create_report_table(self._hsn_headers())
        self.tab_widget.addTab(self.b2b_b2cl_table, 'B2B & B2CL')
        self.tab_widget.addTab(self.b2cs_table, 'B2CS')
        self.tab_widget.addTab(self.hsn_table, 'HSN Summary')
        self.tab_widget.currentChanged.connect(self._on_tab_changed)
        layout.addWidget(self.tab_widget, 1)
        layout.addWidget(self._build_footer_frame())

    def _create_report_table(self, headers):
        """Create a read-only report table with the supplied column headers."""
        table = QTableWidget()
        table.setStyleSheet(report_data_table_style())
        table.setSortingEnabled(False)
        apply_read_only_report_table_selection(table)
        table.setTextElideMode(Qt.ElideNone)
        table.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        table.setColumnCount(len(headers))
        table.setHorizontalHeaderLabels(headers)
        table.itemDoubleClicked.connect(self.on_row_double_clicked)
        apply_compact_report_table_rows(table)
        return table

    def _on_tab_changed(self, tab_index: int) -> None:
        """Track the active tab table for column-width memory."""
        self._activate_tab_table(tab_index)

    def _activate_tab_table(self, tab_index: int) -> None:
        """Restore saved column widths for the selected report tab."""
        table_map = {
            0: ('b2b_b2cl_table', self.b2b_b2cl_table),
            1: ('b2cs_table', self.b2cs_table),
            2: ('hsn_table', self.hsn_table),
        }
        memory_attr, table = table_map.get(tab_index, table_map[0])
        self._ui_memory_active_table_attr = memory_attr
        self._ui_memory_active_table = table
        if hasattr(self, 'settings'):
            self._restore_memory_table(table, memory_attr)

    def load_company(self):
        active = active_company_manager.get_active_company()
        if active:
            self.company_id = active.get('id')
            self.company_state = (active.get('state') or '').strip()

    def generate_report(self):
        if self._loading:
            return
        if not self.company_id:
            QMessageBox.warning(self, 'No Company', 'Please open a company first.')
            return
        from_date = qdate_to_db(self.from_date.date())
        to_date = qdate_to_db(self.to_date.date())
        search = self.search_input.text().strip()
        db_type = getattr(self.db, 'db_type', None)
        db_path = getattr(self.db, 'db_path', None)
        thread = QThread(self)
        worker = GSTSalesReportWorker(db_type, db_path, self.company_id, self.company_state, from_date, to_date, search)
        worker.moveToThread(thread)
        thread.started.connect(worker.run)
        worker.data_ready.connect(self._on_report_ready)
        worker.error.connect(self._on_report_error)
        worker.finished.connect(thread.quit)
        worker.finished.connect(worker.deleteLater)
        thread.finished.connect(thread.deleteLater)
        thread.finished.connect(self._on_report_finished)
        self._report_thread = thread
        self._report_worker = worker
        self._set_loading_state(True)
        thread.start()

    def _set_loading_state(self, is_loading):
        """Disable report controls while the worker calculates."""
        self._loading = is_loading
        self.generate_btn.setEnabled(not is_loading)
        self.export_btn.setEnabled(not is_loading)
        self.export_pdf_btn.setEnabled(not is_loading and PDF_AVAILABLE)
        if is_loading:
            self.status_label.setText('Calculating...')
            self._reset_footer_totals()

    def _on_report_ready(self, report_data):
        """Display worker-calculated report data across all GST sales tabs."""
        self.current_report_data = report_data
        self._display_b2b_b2cl_tab(
            report_data.get('b2b', []),
            report_data.get('b2cl', []),
        )
        self._display_b2cs_tab(report_data.get('b2cs', {}))
        self._display_hsn_tab(report_data.get('hsn', []))
        self._update_footer_totals(report_data)
        self._activate_tab_table(self.tab_widget.currentIndex())
        self.status_label.setText('GST sales report generated successfully.')

    def _on_report_error(self, message):
        """Show worker errors and keep the GUI responsive."""
        QMessageBox.critical(self, 'Error', f'Failed to generate report: {message}')
        self._reset_footer_totals()
        self.status_label.setText('GST sales report generation failed.')

    def _on_report_finished(self):
        """Reset worker references when the thread exits."""
        self._report_thread = None
        self._report_worker = None
        self._set_loading_state(False)

    def _build_footer_frame(self):
        """Create the dark themed GST totals footer."""
        frame = QFrame()
        frame.setStyleSheet(theme.section_panel_frame_style() + '\n            QLabel {\n                background: transparent;\n                border: none;\n            }\n        ')
        footer_layout = QHBoxLayout(frame)
        footer_layout.setContentsMargins(5, 5, 5, 5)
        footer_layout.setSpacing(12)
        self.footer_total_labels = {}
        self.footer_total_titles = {}
        footer_fields = [('Total Invoice Value', 'invoice_value'), ('Total Taxable Value', 'taxable_value'), ('Total CGST', 'cgst'), ('Total SGST', 'sgst'), ('Total IGST', 'igst'), ('Total Cess', 'cess')]
        for title, key in footer_fields:
            metric_label = QLabel()
            metric_label.setTextFormat(Qt.RichText)
            metric_label.setWordWrap(False)
            metric_label.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Fixed)
            metric_label.setStyleSheet('background: transparent; border: none;')
            self._set_footer_metric(metric_label, title, '0.00')
            footer_layout.addWidget(metric_label)
            self.footer_total_labels[key] = metric_label
            self.footer_total_titles[key] = title
        footer_layout.addStretch()
        return frame

    def _footer_metric_html(self, title, value):
        """Return rich footer HTML that keeps a metric title and value together."""
        clean_title = str(title or '').strip()
        if not clean_title.endswith(':'):
            clean_title = f'{clean_title}:'
        return f"<span style='{footer_title_style()}'>{clean_title} </span><span style='{footer_value_style()}'>{value}</span>"

    def _set_footer_metric(self, label, title, value):
        """Update a combined footer metric label without splitting widgets."""
        label.setText(self._footer_metric_html(title, value))

    def _set_footer_total(self, key, value):
        """Set one footer total while preserving its title text."""
        label = self.footer_total_labels.get(key)
        title = self.footer_total_titles.get(key, key)
        if label is not None:
            self._set_footer_metric(label, title, value)

    def _reset_footer_totals(self):
        """Reset footer totals while loading or after errors."""
        if not hasattr(self, 'footer_total_labels'):
            return
        for key in self.footer_total_labels:
            self._set_footer_total(key, '0.00')

    def _update_footer_totals(self, report_data):
        """Calculate footer totals from currently loaded structured report data."""
        totals = {'invoice_value': 0.0, 'taxable_value': 0.0, 'cgst': 0.0, 'sgst': 0.0, 'igst': 0.0, 'cess': 0.0}
        for sale in report_data.get('b2b', []):
            self._add_invoice_totals(totals, sale)
        for sale in report_data.get('b2cl', []):
            self._add_invoice_totals(totals, sale)
        for row in report_data.get('b2cs', {}).values():
            totals['invoice_value'] += self._to_float(row.get('total_value'))
            totals['taxable_value'] += self._to_float(row.get('taxable_value'))
            totals['cgst'] += self._to_float(row.get('cgst'))
            totals['sgst'] += self._to_float(row.get('sgst'))
            totals['igst'] += self._to_float(row.get('igst'))
            totals['cess'] += self._to_float(row.get('cess'))
        for key, value in totals.items():
            self._set_footer_total(key, f'{value:.2f}')

    def _add_invoice_totals(self, totals, sale):
        """Add one invoice row to the footer totals."""
        totals['invoice_value'] += self._to_float(sale.get('grand_total'))
        totals['taxable_value'] += self._to_float(sale.get('taxable_value'))
        totals['cgst'] += self._to_float(sale.get('cgst'))
        totals['sgst'] += self._to_float(sale.get('sgst'))
        totals['igst'] += self._to_float(sale.get('igst'))
        totals['cess'] += self._to_float(sale.get('cess'))

    def _classify_sales(self, sales):
        """Classify sales using the same compliance rules as the worker."""
        b2b_sales = []
        b2cl_sales = []
        b2cs_data = {}
        for sale in sales:
            self._normalise_sale_tax_values(sale)
            gstin = self._gstin(sale)
            state = self._place_of_supply(sale)
            grand_total = self._to_float(sale.get('grand_total'))
            form_of_sale = self._classify_sale(gstin, state, grand_total)
            sale['derived_form_of_sale'] = form_of_sale
            sale['place_of_supply'] = state
            sale['invoice_type'] = 'Regular'
            sale['reverse_charge'] = 'N'
            sale['rate'] = self._get_tax_rate(sale)
            sale['type'] = form_of_sale
            if form_of_sale == 'B2B':
                b2b_sales.append(sale)
            elif form_of_sale == 'B2CL':
                b2cl_sales.append(sale)
            else:
                key = (sale['place_of_supply'], sale['rate'])
                if key not in b2cs_data:
                    b2cs_data[key] = {'type': 'B2CS', 'place_of_supply': sale['place_of_supply'], 'rate': sale['rate'], 'taxable_value': 0.0, 'cgst': 0.0, 'sgst': 0.0, 'igst': 0.0, 'cess': 0.0, 'total_value': 0.0, 'invoice_count': 0, 'sale_ids': []}
                b2cs_data[key]['taxable_value'] += self._to_float(sale.get('taxable_value'))
                b2cs_data[key]['cgst'] += self._to_float(sale.get('cgst'))
                b2cs_data[key]['sgst'] += self._to_float(sale.get('sgst'))
                b2cs_data[key]['igst'] += self._to_float(sale.get('igst'))
                b2cs_data[key]['cess'] += self._to_float(sale.get('cess'))
                b2cs_data[key]['total_value'] += grand_total
                b2cs_data[key]['invoice_count'] += 1
                b2cs_data[key]['sale_ids'].append(sale.get('id'))
        return (b2b_sales, b2cl_sales, b2cs_data)

    def _classify_sale(self, gstin, state, grand_total):
        """Return B2B, B2CL, or B2CS for a sale."""
        return classify_invoice(gstin, state, self.company_state, grand_total)

    def _normalise_sale_tax_values(self, sale):
        taxable = self._to_float(sale.get('taxable_value'))
        tax_total = self._to_float(sale.get('tax_total') or sale.get('item_tax_total'))
        state = self._place_of_supply(sale)
        interstate = is_interstate(state, self.company_state)
        igst, cgst, sgst, cess = tax_totals_for_supply(sale, interstate)
        sale['taxable_value'] = taxable
        sale['cgst'] = cgst
        sale['sgst'] = sgst
        sale['igst'] = igst
        sale['cess'] = cess
        sale['tax_total'] = tax_total or igst + cgst + sgst + cess

    def _gstin(self, sale):
        return normalized_gstin(sale.get('party_gstin') or sale.get('gstin'))

    def _place_of_supply(self, sale):
        return place_of_supply_label(sale.get('party_state') or sale.get('state'), self.company_state)

    def _get_tax_rate(self, sale):
        taxable = self._to_float(sale.get('taxable_value'))
        tax = self._to_float(sale.get('cgst')) + self._to_float(sale.get('sgst')) + self._to_float(sale.get('igst')) + self._to_float(sale.get('cess'))
        if taxable > 0:
            return round(tax / taxable * 100, 2)
        return 0.0

    def _to_float(self, value):
        try:
            return float(value or 0)
        except (TypeError, ValueError):
            return 0.0

    def _b2b_b2cl_headers(self):
        return [
            'GSTIN/UIN', 'Receiver Name', 'Invoice No', 'Invoice Date', 'Invoice Value',
            'Place of Supply', 'Reverse Charge', 'Invoice Type', 'Rate (%)', 'Taxable Value',
            'IGST', 'CGST', 'SGST', 'Cess',
        ]

    def _b2cs_headers(self):
        return [
            'Type', 'Place Of Supply', 'Rate (%)', 'Taxable Value',
            'IGST', 'CGST', 'SGST', 'Cess',
        ]

    def _hsn_headers(self):
        return [
            'HSN/SAC', 'Description', 'UQC', 'Total Quantity', 'Total Value',
            'Taxable Value', 'IGST', 'CGST', 'SGST', 'Cess',
        ]

    def _display_b2b_b2cl_tab(self, b2b_sales, b2cl_sales):
        """Populate the combined B2B and B2CL invoice tab."""
        headers = self._b2b_b2cl_headers()
        combined_sales = list(b2b_sales) + list(b2cl_sales)
        self._populate_table(self.b2b_b2cl_table, combined_sales, headers)

    def _display_b2cs_tab(self, b2cs_data):
        """Populate the consolidated B2CS tab."""
        headers = self._b2cs_headers()
        self._populate_table(self.b2cs_table, list(b2cs_data.values()), headers)

    def _display_hsn_tab(self, hsn_data):
        """Populate the HSN summary tab."""
        headers = self._hsn_headers()
        self._populate_table(self.hsn_table, hsn_data or [], headers)

    def _populate_table(self, table, rows, headers):
        """Fill one report tab table with normalized row values."""
        table.clearContents()
        table.setRowCount(0)
        table.setColumnCount(len(headers))
        table.setHorizontalHeaderLabels(headers)
        if not rows:
            table.insertRow(0)
            item = QTableWidgetItem('No data')
            item.setForeground(QColor(theme.semantic_neutral_hex()))
            item.setFlags(item.flags() & ~Qt.ItemIsEditable)
            table.setItem(0, 0, item)
        else:
            for row_index, data in enumerate(rows):
                table.insertRow(row_index)
                if isinstance(data, dict) and 'sale_ids' in data:
                    metadata = {'type': 'b2cs_group', 'sale_ids': data.get('sale_ids', [])}
                else:
                    metadata = data.get('id') if isinstance(data, dict) else None
                for col, header in enumerate(headers):
                    value = self._value_for_header(data, header)
                    item = self._make_item(value, header)
                    item.setData(Qt.UserRole, metadata)
                    table.setItem(row_index, col, item)
        self._finish_table(table)

    def _value_for_header(self, data, header):
        """Map structured report rows to tab column labels."""
        hsn_total_value = (
            self._to_float(data.get('val'))
            + self._to_float(data.get('iamt'))
            + self._to_float(data.get('camt'))
            + self._to_float(data.get('samt'))
            + self._to_float(data.get('csamt'))
        )
        mapping = {
            'GSTIN/UIN': self._gstin(data),
            'Receiver Name': data.get('party_name', ''),
            'Invoice No': data.get('invoice_number', ''),
            'Invoice Date': format_display_date(data.get('invoice_date', '')),
            'Invoice Value': data.get('grand_total', 0.0),
            'Place of Supply': data.get('place_of_supply') or self._place_of_supply(data),
            'Place Of Supply': data.get('place_of_supply') or self._place_of_supply(data),
            'Reverse Charge': data.get('reverse_charge', 'N'),
            'Invoice Type': data.get('derived_form_of_sale') or data.get('invoice_type', 'Regular'),
            'Rate (%)': data.get('rate', self._get_tax_rate(data)),
            'Taxable Value': data.get('taxable_value', data.get('val', 0.0)),
            'CGST': data.get('cgst', data.get('camt', 0.0)),
            'SGST': data.get('sgst', data.get('samt', 0.0)),
            'IGST': data.get('igst', data.get('iamt', 0.0)),
            'Cess': data.get('cess', data.get('csamt', 0.0)),
            'Type': data.get('type', ''),
            'HSN/SAC': data.get('hsn', ''),
            'Description': data.get('desc', ''),
            'UQC': data.get('uqc', ''),
            'Total Quantity': data.get('qty', 0.0),
            'Total Value': hsn_total_value,
        }
        return mapping.get(header, '')

    def _make_item(self, value, header):
        if header in self.MONEY_KEYS or isinstance(value, float):
            text = f'{self._to_float(value):.2f}'
        else:
            text = str(value if value is not None else '')
        item = QTableWidgetItem(text)
        item.setFlags(item.flags() & ~Qt.ItemIsEditable)
        return item

    def _format_export_value(self, value, header):
        """Format one export cell using the same rules as on-screen report tables."""
        if header in self.MONEY_KEYS or isinstance(value, float):
            return f'{self._to_float(value):.2f}'
        return str(value if value is not None else '')

    def _has_report_data(self) -> bool:
        """Return True when a generated report is available for export."""
        report_data = self.current_report_data or {}
        return bool(
            report_data.get('b2b')
            or report_data.get('b2cl')
            or report_data.get('b2cs')
            or report_data.get('hsn')
        )

    def _export_b2b_rows(self):
        """Return combined B2B and B2CL rows from the cached report payload."""
        report_data = self.current_report_data or {}
        return list(report_data.get('b2b', [])) + list(report_data.get('b2cl', []))

    def _export_b2cs_rows(self):
        """Return consolidated B2CS rows from the cached report payload."""
        report_data = self.current_report_data or {}
        return list(report_data.get('b2cs', {}).values())

    def _export_hsn_rows(self):
        """Return HSN summary rows from the cached report payload."""
        report_data = self.current_report_data or {}
        return list(report_data.get('hsn', []) or [])

    def _build_export_dataframe(self, headers, rows):
        """Build a pandas DataFrame with UI-matching headers and formatted values."""
        import pandas as pd

        export_rows = []
        for row in rows:
            export_rows.append({
                header: self._format_export_value(self._value_for_header(row, header), header)
                for header in headers
            })
        if not export_rows:
            return pd.DataFrame(columns=headers)
        return pd.DataFrame(export_rows, columns=headers)

    def _autosize_worksheet_columns(self, worksheet, dataframe) -> None:
        """Auto-fit worksheet column widths for readable Excel output."""
        from openpyxl.utils import get_column_letter

        for column_index, column_name in enumerate(dataframe.columns, start=1):
            column_values = [str(column_name)]
            if not dataframe.empty:
                column_values.extend(
                    str(value) for value in dataframe[column_name].fillna('').astype(str).tolist()
                )
            max_length = max(len(value) for value in column_values) if column_values else len(str(column_name))
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 10), 50)

    def _build_pdf_table_matrix(self, headers, rows):
        """Build a reportlab table matrix with UI-matching headers and formatted values."""
        matrix = [list(headers)]
        if not rows:
            empty_row = ['No data'] + [''] * (len(headers) - 1)
            matrix.append(empty_row)
            return matrix
        for row in rows:
            matrix.append([
                self._format_export_value(self._value_for_header(row, header), header)
                for header in headers
            ])
        return matrix

    def _make_pdf_styled_table(self, headers, rows, available_width):
        """Create one zebra-striped reportlab table sized for landscape A4."""
        data = self._build_pdf_table_matrix(headers, rows)
        column_count = max(len(headers), 1)
        column_width = available_width / column_count
        table = Table(data, colWidths=[column_width] * column_count, repeatRows=1)
        money_headers = set(self.MONEY_KEYS) | {'Rate (%)'}
        style_commands = [
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#374151')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('FONTSIZE', (0, 1), (-1, -1), 7),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#9CA3AF')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#E5E7EB')]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('LEFTPADDING', (0, 0), (-1, -1), 3),
            ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ]
        for column_index, header in enumerate(headers):
            if header in money_headers:
                style_commands.append(
                    ('ALIGN', (column_index, 1), (column_index, -1), 'RIGHT')
                )
        table.setStyle(TableStyle(style_commands))
        return table

    def export_pdf(self):
        """Export all GST sales sections into one landscape PDF document."""
        if not PDF_AVAILABLE:
            QMessageBox.warning(
                self,
                'Export Failed',
                'reportlab is not installed.\n\nInstall it with:\npip install reportlab',
            )
            return
        if not self._has_report_data():
            QMessageBox.warning(self, 'No Report', 'Please generate the GST sales report first.')
            return

        from_date_db = qdate_to_db(self.from_date.date())
        to_date_db = qdate_to_db(self.to_date.date())
        from_date = qdate_to_display(self.from_date.date())
        to_date = qdate_to_display(self.to_date.date())
        default_name = f'GST_Sales_Report_{from_date_db}_to_{to_date_db}.pdf'
        filepath, _ = QFileDialog.getSaveFileName(
            self,
            'Export GST Sales Report PDF',
            default_name,
            'PDF Files (*.pdf)',
        )
        if not filepath:
            return
        if not filepath.lower().endswith('.pdf'):
            filepath = f'{filepath}.pdf'

        try:
            page_size = landscape(A4)
            document = SimpleDocTemplate(
                filepath,
                pagesize=page_size,
                leftMargin=36,
                rightMargin=36,
                topMargin=36,
                bottomMargin=36,
            )
            available_width = page_size[0] - document.leftMargin - document.rightMargin
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'GstSalesReportTitle',
                parent=styles['Heading1'],
                fontSize=16,
                alignment=1,
                spaceAfter=8,
            )
            subtitle_style = ParagraphStyle(
                'GstSalesReportSubtitle',
                parent=styles['Normal'],
                fontSize=10,
                alignment=1,
                spaceAfter=16,
            )
            section_style = ParagraphStyle(
                'GstSalesReportSection',
                parent=styles['Heading2'],
                fontSize=11,
                spaceBefore=14,
                spaceAfter=8,
            )
            story = [
                Paragraph('Faizan Pro App - GST Sales Report', title_style),
                Paragraph(f'Date Range: {from_date} to {to_date}', subtitle_style),
            ]
            section_specs = [
                ('B2B &amp; B2CL', self._b2b_b2cl_headers(), self._export_b2b_rows()),
                ('B2CS (Consolidated)', self._b2cs_headers(), self._export_b2cs_rows()),
                ('HSN Summary', self._hsn_headers(), self._export_hsn_rows()),
            ]
            for section_title, headers, rows in section_specs:
                story.append(Paragraph(section_title, section_style))
                story.append(self._make_pdf_styled_table(headers, rows, available_width))
                story.append(Spacer(1, 12))
            document.build(story)
            QMessageBox.information(
                self,
                'Export Success',
                f'GST Sales Report PDF exported successfully to:\n{filepath}',
            )
        except Exception as exc:
            QMessageBox.critical(self, 'Export Failed', f'Failed to export PDF: {exc}')

    def _finish_table(self, table):
        """Apply column sizing and a compact standard row height for one report tab."""
        apply_adjustable_table_columns(table)
        apply_compact_report_table_rows(table)

    def _report_table_style(self):
        """Return the shared report table stylesheet for detail dialogs."""
        return self.b2b_b2cl_table.styleSheet()

    def on_row_double_clicked(self, item):
        metadata = item.data(Qt.UserRole)
        if not metadata:
            return
        if isinstance(metadata, dict) and metadata.get('type') == 'b2cs_group':
            self._show_b2cs_group_dialog(metadata.get('sale_ids', []))
        else:
            self._show_sale_details(metadata)

    def _show_b2cs_group_dialog(self, sale_ids):
        """Show dialog listing all bills in a B2CS group."""
        try:
            if not sale_ids:
                QMessageBox.information(self, 'No Bills', 'No bills found in this group.')
                return
            ph = self.db._get_placeholder()
            placeholders = ','.join([ph] * len(sale_ids))
            query = f'\n                SELECT s.*, p.name as party_name, p.gstin as party_gstin, p.state as party_state\n                FROM sales s\n                LEFT JOIN parties p ON s.party_id = p.id\n                WHERE s.id IN ({placeholders})\n                ORDER BY s.invoice_date DESC, s.id DESC\n            '
            sales = self.db.execute_query(query, tuple(sale_ids))
            if not sales:
                QMessageBox.information(self, 'No Bills', 'No bills found in this group.')
                return
            dialog = QDialog(self)
            dialog.setWindowTitle(f'B2CS Group - {len(sales)} Bills')
            dialog.setMinimumSize(1100, 600)
            dialog.setStyleSheet(self._report_table_style() + report_detail_dialog_style())
            layout = QVBoxLayout(dialog)
            header_label = QLabel(f'<b>B2CS Consolidated Group - {len(sales)} Bills</b>')
            header_label.setStyleSheet(page_heading_style(16))
            layout.addWidget(header_label)
            table = QTableWidget()
            apply_read_only_report_table_selection(table)
            headers = ['Invoice No', 'Date', 'Customer', 'State', 'Rate', 'Taxable Value', 'CGST', 'SGST', 'IGST', 'CESS', 'Grand Total']
            table.setColumnCount(len(headers))
            table.setHorizontalHeaderLabels(headers)
            table.setRowCount(len(sales))
            for i, sale in enumerate(sales):
                self._normalise_sale_tax_values(sale)
                rate = self._get_tax_rate(sale)
                values = [sale.get('invoice_number', ''), format_display_date(sale.get('invoice_date', '')), sale.get('party_name', ''), sale.get('party_state', '') or sale.get('state', ''), f'{rate:.2f}', f"{self._to_float(sale.get('taxable_value')):.2f}", f"{self._to_float(sale.get('cgst')):.2f}", f"{self._to_float(sale.get('sgst')):.2f}", f"{self._to_float(sale.get('igst')):.2f}", f"{self._to_float(sale.get('cess')):.2f}", f"{self._to_float(sale.get('grand_total')):.2f}"]
                for col, value in enumerate(values):
                    item = QTableWidgetItem(str(value))
                    item.setData(Qt.UserRole, sale.get('id'))
                    item.setFlags(item.flags() & ~Qt.ItemIsEditable)
                    table.setItem(i, col, item)
            apply_adjustable_table_columns(table)
            apply_compact_report_table_rows(table)
            table.itemDoubleClicked.connect(lambda item: self._show_sale_details(item.data(Qt.UserRole)))
            layout.addWidget(table)
            close_btn = QPushButton('Close')
            close_btn.clicked.connect(dialog.accept)
            layout.addWidget(close_btn, alignment=Qt.AlignRight)
            dialog.exec()
        except Exception as e:
            QMessageBox.critical(self, 'Error', f'Failed to load B2CS group: {e}')

    def _show_sale_details(self, sale_id):
        try:
            ph = self.db._get_placeholder()
            query = f'\n                SELECT s.*, p.name as party_name, p.gstin as party_gstin, p.state as party_state\n                FROM sales s\n                LEFT JOIN parties p ON s.party_id = p.id\n                WHERE s.id = {ph}\n            '
            result = self.db.execute_query(query, (sale_id,))
            if not result:
                QMessageBox.information(self, 'Not Found', 'Sale not found.')
                return
            sale = result[0]
            self._normalise_sale_tax_values(sale)
            items_query = f'\n                SELECT si.*, prod.name as product_name, prod.hsn\n                FROM sales_items si\n                LEFT JOIN products prod ON si.product_id = prod.id\n                WHERE si.sale_id = {ph}\n                ORDER BY si.sl_no\n            '
            items = self.db.execute_query(items_query, (sale_id,))
            self._show_invoice_details_dialog(sale, items)
        except Exception as e:
            QMessageBox.critical(self, 'Error', f'Failed to load sale details: {e}')

    def _show_invoice_details_dialog(self, sale, items):
        """Show full invoice details with all required fields."""
        dialog = QDialog(self)
        dialog.setWindowTitle(f"Invoice Details - {sale.get('invoice_number', '')}")
        dialog.setMinimumSize(1200, 700)
        dialog.setStyleSheet(self._report_table_style() + report_detail_dialog_style())
        layout = QVBoxLayout(dialog)
        form_of_sale = sale.get('form_of_sale') or sale.get('derived_form_of_sale', '')
        header_text = f"<b>Invoice No:</b> {sale.get('invoice_number', '')} &nbsp;&nbsp; <b>Date:</b> {format_display_date(sale.get('invoice_date', ''))} &nbsp;&nbsp; <b>Customer:</b> {sale.get('party_name', '')} &nbsp;&nbsp; <b>GSTIN:</b> {sale.get('party_gstin') or sale.get('gstin') or ''} &nbsp;&nbsp; <b>State:</b> {sale.get('party_state') or sale.get('state') or ''} &nbsp;&nbsp; <b>Form of Sale:</b> {form_of_sale} &nbsp;&nbsp; <b>Nature:</b> {sale.get('nature', '')}"
        header_label = QLabel(header_text)
        header_label.setWordWrap(True)
        header_label.setStyleSheet(page_heading_style(14))
        layout.addWidget(header_label)
        summary_grid = QGridLayout()
        summary_grid.setSpacing(10)
        colors = theme._theme_colors()
        financial_fields = [('Bill Value:', self._to_float(sale.get('grand_total'))), ('Taxable Value:', self._to_float(sale.get('taxable_value'))), ('CGST:', self._to_float(sale.get('cgst'))), ('SGST:', self._to_float(sale.get('sgst'))), ('IGST:', self._to_float(sale.get('igst'))), ('CESS:', self._to_float(sale.get('cess'))), ('Discount:', self._to_float(sale.get('discount_total'))), ('Round Off:', self._to_float(sale.get('round_off'))), ('Amount Received:', self._to_float(sale.get('amount_received')))]
        for i, (label, value) in enumerate(financial_fields):
            lbl = QLabel(label)
            lbl.setStyleSheet(report_summary_label_style())
            val = QLabel(f'{value:.2f}')
            val.setStyleSheet(f"color: {colors['input_text']}; font-size: 13px; background: transparent; border: none;")
            summary_grid.addWidget(lbl, i // 3, i % 3 * 2)
            summary_grid.addWidget(val, i // 3, i % 3 * 2 + 1)
        summary_widget = QWidget()
        summary_widget.setLayout(summary_grid)
        summary_widget.setStyleSheet(report_filter_frame_style())
        layout.addWidget(summary_widget)
        items_label = QLabel('<b>Items</b>')
        items_label.setStyleSheet(page_heading_style(14))
        layout.addWidget(items_label)
        items_table = QTableWidget()
        apply_read_only_report_table_selection(items_table)
        item_headers = ['SL', 'Product', 'HSN', 'Rate', 'Qty', 'Gross', 'Discount', 'Tax %', 'Tax Amount', 'Total']
        items_table.setColumnCount(len(item_headers))
        items_table.setHorizontalHeaderLabels(item_headers)
        items_table.setRowCount(len(items))
        for i, item in enumerate(items):
            qty = self._to_float(item.get('quantity', 0))
            rate = self._to_float(item.get('rate', 0))
            gross = qty * rate
            discount = self._to_float(item.get('discount', 0))
            tax_amount = self._to_float(item.get('tax_amount', 0))
            total = self._to_float(item.get('grand_total', 0))
            taxable = self._to_float(item.get('net_value', 0))
            tax_pct = round(tax_amount / taxable * 100, 2) if taxable > 0 else 0.0
            values = [item.get('sl_no', ''), item.get('product_name', ''), item.get('hsn', ''), f'{rate:.2f}', f'{qty:.2f}', f'{gross:.2f}', f'{discount:.2f}', f'{tax_pct:.2f}', f'{tax_amount:.2f}', f'{total:.2f}']
            for col, value in enumerate(values):
                table_item = QTableWidgetItem(str(value))
                table_item.setFlags(table_item.flags() & ~Qt.ItemIsEditable)
                items_table.setItem(i, col, table_item)
        apply_adjustable_table_columns(items_table)
        apply_compact_report_table_rows(items_table)
        layout.addWidget(items_table)
        close_btn = QPushButton('Close')
        close_btn.clicked.connect(dialog.accept)
        layout.addWidget(close_btn, alignment=Qt.AlignRight)
        dialog.exec()

    def export_excel(self):
        """Export all GST sales tabs into one Excel workbook with three worksheets."""
        if not self._has_report_data():
            QMessageBox.warning(self, 'No Report', 'Please generate the GST sales report first.')
            return

        from_date = qdate_to_db(self.from_date.date())
        to_date = qdate_to_db(self.to_date.date())
        default_name = f'GST_Sales_Report_{from_date}_to_{to_date}.xlsx'
        filepath, _ = QFileDialog.getSaveFileName(
            self,
            'Export GST Sales Report',
            default_name,
            'Excel Files (*.xlsx)',
        )
        if not filepath:
            return
        if not filepath.lower().endswith('.xlsx'):
            filepath = f'{filepath}.xlsx'

        try:
            import pandas as pd
        except ImportError:
            QMessageBox.warning(
                self,
                'Export Failed',
                'pandas is not installed.\n\nInstall it with:\npip install pandas openpyxl',
            )
            return

        try:
            import openpyxl  # noqa: F401
        except ImportError:
            QMessageBox.warning(
                self,
                'Export Failed',
                'openpyxl is not installed.\n\nInstall it with:\npip install pandas openpyxl',
            )
            return

        try:
            sheet_specs = [
                ('b2b', self._b2b_b2cl_headers(), self._export_b2b_rows()),
                ('b2cs', self._b2cs_headers(), self._export_b2cs_rows()),
                ('hsn', self._hsn_headers(), self._export_hsn_rows()),
            ]
            with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
                for sheet_name, headers, rows in sheet_specs:
                    dataframe = self._build_export_dataframe(headers, rows)
                    dataframe.to_excel(writer, sheet_name=sheet_name, index=False)
                    self._autosize_worksheet_columns(writer.sheets[sheet_name], dataframe)
            QMessageBox.information(
                self,
                'Export Success',
                f'GST Sales Report exported successfully to:\n{filepath}',
            )
        except Exception as exc:
            QMessageBox.critical(self, 'Export Failed', f'Failed to export: {exc}')

    def refresh_theme(self) -> None:
        """Re-apply theme-aware styles after a global theme change."""
        self.setStyleSheet(report_compound_entry_page_style())
        if hasattr(self, 'title_label'):
            self.title_label.setStyleSheet(page_heading_style(18))
        if hasattr(self, 'filter_frame'):
            self.filter_frame.setStyleSheet(report_filter_frame_style())
        if hasattr(self, 'status_label'):
            self.status_label.setStyleSheet(report_status_label_style())
        for table in (getattr(self, 'b2b_b2cl_table', None), getattr(self, 'b2cs_table', None), getattr(self, 'hsn_table', None)):
            if table is not None:
                table.setStyleSheet(report_data_table_style())